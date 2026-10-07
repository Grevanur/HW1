from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'Gowtham_Revanur_Implementation.docx'

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def text(doc, value, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(value)
    run.bold = bold
    return p

def heading(doc, value, level=1):
    p = doc.add_heading(value, level=level)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    return p

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.75)
section.bottom_margin = Inches(.75)
section.left_margin = Inches(.8)
section.right_margin = Inches(.8)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10.5)
styles['Title'].font.name = 'Aptos Display'
styles['Title'].font.size = Pt(22)
styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
for style_name in ['Heading 1', 'Heading 2', 'Heading 3']:
    styles[style_name].font.name = 'Aptos Display'
    styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('Calculator Studio Implementation Rationale')
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run('CSC 6370 Mobile Application Development  |  Homework 1').italic = True

table = doc.add_table(rows=2, cols=2)
table.style = 'Table Grid'
fields = [('Student', 'Gowtham Revanur'), ('Student ID', '002574540'),
          ('Repository', 'https://github.com/Grevanur/HW1'), ('Date', 'October 7, 2026')]
for idx, (label, value) in enumerate(fields):
    cell = table.cell(idx // 2, (idx % 2))
    cell.text = ''
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.add_run(label + '\n').bold = True
    p.add_run(value)
for row in table.rows:
    for cell in row.cells:
        shade(cell, 'F2F5FA')

heading(doc, 'Overview')
text(doc, 'Calculator Studio is a Flutter calculator designed around a deliberate two-operand interaction model. The completed project satisfies the common calculator requirements, the required graduate decimal support, and three graduate advanced features: calculation history, chained operations with a running total, and advanced error handling. The implementation keeps calculator state local to one StatefulWidget so every tap causes a clear, inspectable state transition.')

heading(doc, 'Selected Graduate Features and Evidence')
feature_table = doc.add_table(rows=1, cols=3)
feature_table.style = 'Table Grid'
headers = ['Requirement', 'Implementation evidence', 'Verification evidence']
for cell, value in zip(feature_table.rows[0].cells, headers):
    cell.text = value
    shade(cell, '1F4E78')
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
    set_cell_margins(cell)
rows = [
 ('Decimal support', 'The decimal key appends one point only when the display has none; decimal values flow through all four operations.', 'Automated tests cover decimal addition, subtraction, multiplication, and division.'),
 ('Calculation history', 'Each completed operation is stored as a Calculation record. A scrollable list displays expression and result; tapping an entry reuses its result. The app bar clears history.', 'Manual test: complete 1.5 + 2.25 = 3.75, tap 3.75 in history, then continue a new calculation.'),
 ('Multiple operations', 'When another operator is selected, the pending operation is completed and saved before the new operator is stored. The status line announces the running total.', 'Manual test: 8 + 2 × 3 produces 30 using documented left-to-right evaluation.'),
 ('Advanced error handling', 'The app gives recoverable messages for incomplete expressions, division by zero, and non-finite or excessively large results. AC restores a known-good state.', 'Automated test rejects division by zero; manual checks exercise incomplete input and recovery with AC.'),
]
for values in rows:
    cells = feature_table.add_row().cells
    for cell, value in zip(cells, values):
        cell.text = value
        set_cell_margins(cell)

heading(doc, 'Design and Evaluation Strategy')
text(doc, 'I selected left-to-right evaluation for chained operations. In a compact two-operand calculator, the display and running-total message make each intermediate operation visible. For example, entering 8 + 2 × 3 first resolves 8 + 2 to 10 when × is pressed, then resolves 10 × 3 to 30. The alternative is precedence-aware expression parsing. That approach matches a scientific calculator but needs a token model, parser, and a clearer expression display to avoid surprising users. The left-to-right choice keeps the state model small and makes the behavior explainable in the interface and in tests.')
text(doc, 'The central state consists of display text, stored first operand, pending operator, a replacement flag, a user message, and history. A result does not have a second, independently stored display value; instead, the display remains the single source of truth. This avoids drift between result text and what the user sees. Buttons have semantic labels, result and error messages use a live region, and the layout moves the history panel below the keypad on narrow screens.')

heading(doc, 'Required Reflection Questions')
answers = [
 ('01 State and architecture', 'The values that must persist across taps are the display text, the stored operand, pending operator, replacement flag, and history. The display is stored as text because it represents partial input such as 0. or a number the user is actively editing. The arithmetic conversion happens only when an operator or equals requires a numeric value.'),
 ('02 Correctness under pressure', 'I tested ordinary calculations, decimal arithmetic, chained operations, incomplete expressions, and division by zero. The key failure risk is accepting a visually plausible but invalid sequence. The app prevents a second decimal point and asks the user to supply an operator and second value before equals can complete a calculation.'),
 ('03 Accessibility is behavior', 'Accessibility changes interaction, not only appearance. Every key receives a specific semantic label, including All clear, Delete last digit, and Decimal point. The display and status text are live regions so TalkBack can announce a new result or recoverable error without requiring the user to search the screen.'),
 ('04 Performance without visual loss', 'The calculator performs constant-time arithmetic and maintains a small in-memory history. I used a single stateful screen and standard Material controls rather than adding packages or rebuilding unrelated state. The history list is lazy through ListView.builder, which keeps the scrollable feature efficient as entries grow.'),
 ('05 AI output under review', 'AI suggestions were treated as starting points, not evidence. I accepted the recommendation to test normal, boundary, and invalid paths, then verified the arithmetic behavior with flutter test and the actual source code. I rejected any claim that a visible control proves a feature works; each selected feature is tied to a code path and test or manual scenario in this document.'),
 ('06 AI advice trade offs and maintainability', 'The maintainable approach is a narrow CalculatorEngine for arithmetic plus a page that owns interaction state. A more feature-heavy expression parser would be justified only if precedence-aware calculation became a requirement. Keeping one display source of truth and a Calculation value type makes future changes such as persistence or a scientific keypad easier to reason about.'),
 ('Graduate extension Evaluation strategy', 'The chained-operation strategy is intentional left-to-right evaluation. It favors transparency in a basic calculator over conventional precedence. The running-total message and history show the actual sequence, so a user can verify why 8 + 2 × 3 becomes 30. If conventional precedence were required, I would replace the current immediate evaluation with tokenization and explicit parsing, then expand the display and tests to show the full expression.'),
]
for question, answer in answers:
    heading(doc, question, 2)
    text(doc, answer)

heading(doc, 'AI Agent Test Drive Comparison')
text(doc, 'I sent the two unchanged course prompts to OpenAI Codex and Google Gemini on October 7, 2026. I compared their advice against the implemented calculator and the automated test suite. Gemini is identified as a second agent; its response is reported as advice, not as proof that the implementation works.')
prompts = [
 ('Prompt 01 Bug Hunt', 'For a two-operand calculator with +, −, ×, and ÷, propose six test cases with exact inputs and expected outcomes. Include normal, boundary, and invalid sequences. Mark which cases require optional error handling or graduate decimal support. Do not write code.', 'Test normal addition (2 + 3 = 5), decimal multiplication (1.5 × 2 = 3.0; graduate decimal support), subtraction yielding zero (4 − 4 = 0), division (9 ÷ 3 = 3), division by zero (7 ÷ 0 produces a recoverable error; advanced error handling), and incomplete input (5 + = produces a recoverable message; advanced error handling).', 'Gemini proposed basic addition, chained operations, a maximum-input boundary, division by zero, operator overwrite, and decimal precision. Its useful overlap was division by zero and decimal coverage. I accepted those cases and verified the division-by-zero claim against CalculatorEngine.apply in lib/main.dart and the passing flutter test. I rejected Gemini’s assumption that this calculator supports operator overwrite; the implemented interface intentionally evaluates the existing operation when another operator is selected.'),
 ('Prompt 02 State Design', 'A calculator stores displayText, firstOperand, pendingOperator, resultText, and isError. Which values need to be stored, which can be derived, and what bug could happen if resultText and displayText drift apart? Suggest one test that catches it.', 'Store displayText, firstOperand, pendingOperator, and an error indicator. Derive resultText from displayText rather than storing both independently. Otherwise a new digit could update displayText while resultText remains stale. Test by calculating 2 + 3 =, entering 4, and verifying the display announces and shows 4 rather than the old result.', 'Gemini also recommended storing displayText, firstOperand, pendingOperator, and isError while deriving resultText. It described a stale-display bug after clearing or continuing from a result and proposed a reset-and-next-operation test. I accepted the single-source-of-truth recommendation: this implementation has _display but no separate resultText. I verified the claim by reviewing _equals and _enterDigit in lib/main.dart; after equals, entering a digit replaces _display instead of retaining a second stale result value.'),
]
for label, prompt, response, comparison in prompts:
    heading(doc, label, 2)
    text(doc, 'Unchanged prompt: ' + prompt, bold=True)
    text(doc, 'OpenAI Codex response summary: ' + response)
    text(doc, 'Google Gemini response and comparison: ' + comparison)

doc.add_page_break()
heading(doc, 'Testing Record')
test_table = doc.add_table(rows=1, cols=3)
test_table.style = 'Table Grid'
for cell, value in zip(test_table.rows[0].cells, ['Check', 'Expected result', 'Status']):
    cell.text = value
    shade(cell, '1F4E78')
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
    set_cell_margins(cell)
for values in [
 ('flutter analyze', 'No static-analysis issues', 'Completed'),
 ('flutter test', 'Three tests pass, including all four decimal operations and division by zero', 'Completed'),
 ('flutter build apk --release', 'Release APK generated', 'Completed'),
 ('Android device or emulator installation', 'Install and exercise the release APK before upload', 'Student action required'),
 ('Repository visibility', 'Instructor can access complete source and README', 'Student action required'),
]:
    cells = test_table.add_row().cells
    for cell, value in zip(cells, values):
        cell.text = value
        set_cell_margins(cell)

text(doc, 'AI-use disclosure: I used OpenAI Codex for implementation assistance and test planning and Google Gemini for the required independent-agent comparison. I reviewed the suggestions and verified implementation claims using flutter analyze, flutter test, a release build, and source-code inspection.', bold=True)
doc.save(OUT)
print(OUT)
